"""Turn a file into a list of (page_number, text). Supports pdf, docx, txt, md."""

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader


@dataclass
class Extracted:
    pages: list[tuple[int, str]] = field(default_factory=list)
    warning: str | None = None

    @property
    def text(self) -> str:
        return "\n\n".join(t for _, t in self.pages)


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1252"):  # encoding fallback chain
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def extract(path: Path) -> Extracted:
    ext = path.suffix.lower().lstrip(".")
    if ext in ("txt", "md"):
        pages = [(1, _read_text(path))]
    elif ext == "docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:  # tables hold a lot of invoice data
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        pages = [(1, "\n\n".join(parts))]  # DOCX has no fixed pages
    elif ext == "pdf":
        reader = PdfReader(str(path))
        pages = [(i + 1, (pg.extract_text() or "")) for i, pg in enumerate(reader.pages)]
    else:
        raise ValueError(f"Unsupported extension: {ext}")
    pages = [(n, t.replace("\r\n", "\n").replace("\r", "\n")) for n, t in pages if t.strip()]
    warning = None if pages else "no extractable text (scanned PDF or empty file?)"
    return Extracted(pages=pages, warning=warning)


def page_count(path: Path) -> int:
    if path.suffix.lower() == ".pdf":
        try:
            return len(PdfReader(str(path)).pages)
        except Exception:  # noqa: BLE001
            return 0
    return 1
